from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT.parent / "SentinelAI_AWS_Agent_Enterprise_Presentation.pptx"
DIAGRAMS = ROOT / "diagrams"


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullets_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(22)


def add_image_slide(prs: Presentation, title: str, image_name: str, note: str | None = None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    image_path = DIAGRAMS / image_name
    if image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(0.6), Inches(1.2), width=Inches(12.1), height=Inches(5.8))
    else:
        box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.0), Inches(2.0))
        tf = box.text_frame
        tf.text = f"Diagram not found: {image_name}"
        tf.paragraphs[0].font.size = Pt(28)
    if note:
        box = slide.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(12.0), Inches(0.6))
        tf = box.text_frame
        tf.text = note
        tf.paragraphs[0].font.size = Pt(16)


def add_black_video_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0, 0, 0)

    box = slide.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(11.5), Inches(1.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = "Live Incident Walkthrough Video"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

    box2 = slide.shapes.add_textbox(Inches(2.0), Inches(4.3), Inches(9.5), Inches(1.0))
    tf2 = box2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "Embed your demo recording on this slide"
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(210, 210, 210)
    p2.alignment = PP_ALIGN.CENTER


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "SentinelAI AWS Agent",
        "Autonomous Incident Diagnosis and Remediation Co-Pilot for Cloud Operations",
    )

    add_bullets_slide(
        prs,
        "Business Use Case and What We Are Solving",
        [
            "Cloud operations teams face alert fatigue and slow incident resolution.",
            "Root cause analysis is often manual, fragmented, and person-dependent.",
            "SentinelAI automates diagnosis, recommends remediation, and preserves operational memory.",
            "Outcome: faster MTTR, consistent response quality, and lower outage cost.",
        ],
    )

    add_bullets_slide(
        prs,
        "Real-World Use Cases and Latest Global Incident Patterns",
        [
            "CrowdStrike outage (2024) exposed global dependence on rapid diagnosis and coordinated rollback.",
            "Cloud and SaaS incidents continue to show delayed root-cause convergence across teams.",
            "Enterprises need decision-centric operations, not only alert-centric monitoring.",
            "SentinelAI provides explainable investigation and guided remediation for high-severity events.",
        ],
    )

    add_image_slide(
        prs,
        "High-Level Diagram",
        "system_architecture.png",
        "Event-driven flow from CloudWatch alarm to multi-agent AI investigation.",
    )

    add_image_slide(
        prs,
        "Architectural Diagram",
        "aws_service_integration.png",
        "AWS-native architecture with Bedrock, Knowledge Base, API, and memory components.",
    )

    add_black_video_slide(prs)

    add_bullets_slide(
        prs,
        "Business Value in Terms of Cost",
        [
            "Illustrative model: 40 P1/P2 incidents per month, 30% MTTR reduction.",
            "If outage cost is $8,000/hour, potential avoided impact is about $192K per month.",
            "Annualized impact can exceed $2.3M depending on incident mix and service criticality.",
            "Additional gains: less on-call burnout, better auditability, and stronger customer trust.",
        ],
    )

    add_bullets_slide(
        prs,
        "Future Scope and Enhancements",
        [
            "ITSM integrations for automated ticket enrichment and assignment.",
            "Guardrailed auto-remediation using approved runbooks and policy controls.",
            "Predictive incident prevention and risk scoring from telemetry trends.",
            "Multi-cloud extension and security operations workflow convergence.",
        ],
    )

    add_bullets_slide(
        prs,
        "Summary - Striking Close",
        [
            "SentinelAI shifts operations from reactive firefighting to autonomous resilience.",
            "Faster diagnosis, safer remediation, and measurable business value at enterprise scale.",
            "Every critical minute matters: SentinelAI turns minutes into certainty.",
        ],
    )

    prs.save(str(OUT))
    return OUT


if __name__ == "__main__":
    out_path = build()
    print(out_path)
